"""Create styled PowerPoint files, including speaker notes."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

_THEMES: dict[str, dict[str, str]] = {
    "Indigo Modern": {
        "background": "F5F6FF",
        "accent": "4F46E5",
        "text": "1F2340",
        "muted": "626987",
    },
    "Clean White": {
        "background": "FFFFFF",
        "accent": "0F766E",
        "text": "172B2A",
        "muted": "526663",
    },
    "Dark Elegant": {
        "background": "171923",
        "accent": "9F7AEA",
        "text": "F7FAFC",
        "muted": "CBD5E0",
    },
    "Playful": {
        "background": "FFF7ED",
        "accent": "EA580C",
        "text": "292524",
        "muted": "57534E",
    },
}


def _rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def build_pptx(
    *,
    topic: str,
    theme: str,
    slides: Sequence[Mapping[str, Any]],
    speaker_notes: bool,
) -> bytes:
    """Return a complete PPTX; input slide data matches the /api/ppt JSON shape."""
    if theme not in _THEMES:
        raise ValueError(f"Unsupported presentation theme: {theme}")
    if not slides:
        raise ValueError("At least one slide is required.")

    colors = _THEMES[theme]
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    for index, item in enumerate(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(colors["background"])

        accent_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            presentation.slide_width,
            Inches(0.18),
        )
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = _rgb(colors["accent"])
        accent_bar.line.fill.background()

        title = slide.shapes.add_textbox(
            Inches(0.75), Inches(0.55), Inches(11.8), Inches(0.95)
        )
        title_frame = title.text_frame
        title_frame.word_wrap = True
        title_frame.text = str(item.get("title") or topic)
        title_paragraph = title_frame.paragraphs[0]
        title_paragraph.font.name = "Aptos Display"
        title_paragraph.font.size = Pt(30 if index == 0 else 26)
        title_paragraph.font.bold = True
        title_paragraph.font.color.rgb = _rgb(colors["accent"])
        if index == 0:
            title_paragraph.alignment = PP_ALIGN.CENTER

        bullets = item.get("bullets") or []
        body = slide.shapes.add_textbox(
            Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.75)
        )
        body_frame = body.text_frame
        body_frame.clear()
        body_frame.word_wrap = True
        for bullet_index, bullet in enumerate(bullets):
            paragraph = (
                body_frame.paragraphs[0]
                if bullet_index == 0
                else body_frame.add_paragraph()
            )
            paragraph.text = str(bullet)
            paragraph.level = 0
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(22 if len(bullets) <= 3 else 18)
            paragraph.font.color.rgb = _rgb(colors["text"])
            paragraph.space_after = Pt(16)

        diagram = item.get("diagram")
        if diagram:
            diagram_box = slide.shapes.add_textbox(
                Inches(1.0), Inches(6.55), Inches(11.3), Inches(0.5)
            )
            paragraph = diagram_box.text_frame.paragraphs[0]
            paragraph.text = f"Diagram source: {diagram}"
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(8)
            paragraph.font.color.rgb = _rgb(colors["muted"])

        if speaker_notes:
            notes = str(item.get("notes") or "")
            if diagram:
                notes = f"{notes}\n\nMermaid diagram source:\n{diagram}".strip()
            slide.notes_slide.notes_text_frame.text = notes

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
