import json
import re

import pytest
from sqlalchemy import select
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app
from app.models import (
    ChatMessage,
    ChatSession,
    PracticeAttempt,
    PracticeEvaluation,
    PracticePaper,
    VivaSession,
)
from app.services.documents import DocumentPage
from app.services.llm import OmniRouteError


@pytest.fixture
def client():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    test_session = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

    def override_db():
        session = test_session()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()
    engine.dispose()


def test_model_health_and_chat_router_contracts(client, monkeypatch):
    async def fake_models():
        return True, ["teacher-v1", "reasoner-v1"]

    async def fake_stream(messages, task="teacher", temperature=0.4):
        yield "Good "
        yield "morning."

    monkeypatch.setattr("app.main.list_models", fake_models)
    monkeypatch.setattr("app.main.check_models", lambda: fake_models())
    monkeypatch.setattr("app.routers.chat.stream", fake_stream)
    models = client.get("/api/models")
    assert models.status_code == 200
    assert {item["role"] for item in models.json()["models"]} == {
        "teacher", "reasoning", "multimodal", "json"
    }
    response = client.post(
        "/api/chat",
        json={"question": "Explain gravity", "level": "c9-10", "style": "Simple"},
    )
    assert response.status_code == 200
    assert response.text.rstrip().endswith("data: [DONE]")
    assert '"token": "Good "' in response.text


def test_chat_request_uses_current_session_recent_history_and_level_style(client, monkeypatch, capsys):
    from app.models import ChatMessage, ChatSession

    session_id = "synthetic-chat-c1-5"
    other_id = "synthetic-chat-other"
    dependency = app.dependency_overrides[get_db]()
    db = next(dependency)
    db.add_all([ChatSession(id=session_id), ChatSession(id=other_id)])
    for turn in range(11):
        db.add_all(
            [
                ChatMessage(
                    session_id=session_id,
                    role="user",
                    content=f"CURRENT TEST QUESTION {turn}",
                ),
                ChatMessage(
                    session_id=session_id,
                    role="assistant",
                    content=f"CURRENT TEST ANSWER {turn}",
                ),
            ]
        )
    db.add(
        ChatMessage(
            session_id=other_id,
            role="user",
            content="DO NOT INCLUDE OTHER SESSION SENTINEL",
        )
    )
    db.commit()

    captured = []

    async def fake_stream(messages, task="teacher", temperature=0.4):
        captured.extend(messages)
        yield "A synthetic answer."

    monkeypatch.setattr("app.routers.chat.stream", fake_stream)
    try:
        response = client.post(
            "/api/chat",
            json={
                "question": "Why is the sky blue?",
                "level": "c1-5",
                "style": "Exam Answer",
                "subject": "Science",
                "session_id": session_id,
            },
        )
        assert response.status_code == 200
    finally:
        dependency.close()

    printed_messages = json.dumps(captured, ensure_ascii=False, indent=2)
    print("Captured /api/chat model messages:")
    print(printed_messages)
    system_content = captured[0]["content"]
    history = captured[1:-1]
    assert "Classes 1–5" in system_content
    assert "Definition, Explanation, Example" in system_content
    assert "latest question directly and stay on that topic" in system_content
    assert len(history) == 20
    assert history[0]["content"] == "CURRENT TEST QUESTION 1"
    assert history[-1]["content"] == "CURRENT TEST ANSWER 10"
    assert captured[-1] == {"role": "user", "content": "Why is the sky blue?"}
    assert all("DO NOT INCLUDE OTHER SESSION" not in item["content"] for item in captured)
    assert "CURRENT TEST QUESTION 0" not in printed_messages
    assert "current-chat-other" not in printed_messages


def test_solver_router_preserves_paper_question_shape(client, monkeypatch):
    from app.routers import solve

    async def fake_read_validated_upload(file):
        return file.filename, await file.read()

    async def fake_extract_document(data, filename, content_type):
        return [DocumentPage(1, "Question paper fixture for the solver contract test.")]

    def fake_json_call(messages, response_model, task="json"):
        if response_model is solve.ExtractedQuestionSet:
            return response_model(questions=[{"number": 1, "text": "Define gravity.", "marks": 2}])
        return response_model(
            answer="Gravity attracts masses.",
            given=["Masses attract one another."],
            steps=["The force acts between objects with mass."],
            hints=["Think about why objects fall.", "Gravity is the attractive force."],
        )

    async def async_fake_json_call(*args, **kwargs):
        return fake_json_call(*args, **kwargs)

    monkeypatch.setattr(solve, "read_validated_upload", fake_read_validated_upload)
    monkeypatch.setattr(solve, "extract_document", fake_extract_document)
    monkeypatch.setattr(solve, "json_call", async_fake_json_call)
    response = client.post(
        "/api/solve",
        files={"file": ("paper.pdf", b"uploaded test bytes", "application/pdf")},
        data={"level": "c9-10", "subject": "Science", "mode": "exam", "style": "Exam"},
    )
    assert response.status_code == 200
    question = response.json()[0]
    assert set(question) == {"id", "n", "text", "marks", "status", "hints", "solution"}
    assert set(question["solution"]) == {"given", "steps", "therefore", "answer", "blocks"}


def test_documents_upload_list_ask_generate_and_delete_contracts(client, monkeypatch):
    from app.routers import docs

    async def fake_extract(data, filename, content_type):
        return [DocumentPage(12, "Chapter 1\nGravity attracts objects.")]

    async def fake_answer(pages, question, level, estimate):
        return "Gravity attracts masses. [Page 12]", [pages[0]]

    async def fake_generate(pages, task, estimate):
        return {"kind": task, "md": "Generated notes."}

    monkeypatch.setattr(docs, "extract_document", fake_extract)
    monkeypatch.setattr(docs, "answer_document", fake_answer)
    monkeypatch.setattr(docs, "generate_document_result", fake_generate)
    uploaded = client.post(
        "/api/docs/upload",
        files={"file": ("../../chapter.pdf", b"%PDF-1.7 valid-for-mock", "application/pdf")},
    )
    assert uploaded.status_code == 200
    document = uploaded.json()
    assert document["title"] == "chapter.pdf"
    assert set(document) == {"id", "title", "pages", "uploadedAt", "sections"}
    listing = client.get("/api/docs")
    assert [item["id"] for item in listing.json()] == [document["id"]]
    answer = client.post(
        f"/api/docs/{document['id']}/ask",
        json={"question": "What is gravity?", "level": "c9-10"},
    )
    assert answer.json() == {
        "text": "Gravity attracts masses. [Page 12]",
        "sources": [{"page": 12, "quote": "Chapter 1\nGravity attracts objects."}],
    }
    assert client.post(
        f"/api/docs/{document['id']}/generate", json={"task": "notes"}
    ).json() == {"kind": "notes", "md": "Generated notes."}
    assert client.delete(f"/api/docs/{document['id']}").status_code == 204


def test_viva_session_answer_and_report_shapes(client, monkeypatch):
    from app.routers import viva

    async def fake_json_call(messages, response_model, task="json", **kwargs):
        if response_model is viva.VivaQuestionSet:
            return response_model(
                questions=[
                    {
                        "question": f"Question {index}",
                        "topic": "Normalization",
                        "keywords": ["data", "anomaly"],
                        "explanation": "Normalization reduces data duplication.",
                        "followUp": "What is an update anomaly?",
                    }
                    for index in range(5)
                ]
            )
        return response_model(
            label="Good",
            score=1,
            explanation="Clear answer — you covered the main idea.",
            ideal="Normalization reduces duplication.",
        )

    monkeypatch.setattr(viva, "json_call", fake_json_call)
    started = client.post(
        "/api/viva/start",
        json={"subject": "DBMS", "topic": "Normalization", "level": "grad", "count": 5, "adaptive": True},
    )
    assert started.status_code == 200, started.text
    question = started.json()[0]
    assert set(question) == {"id", "question", "topic", "keywords", "explanation", "followUp"}
    answer = {
        "question": question,
        "answer": "Normalization removes duplicate data and update anomalies.",
        "level": "grad",
    }
    evaluated = client.post("/api/viva/answer", json=answer)
    assert evaluated.status_code == 200
    assert evaluated.json() == {
        "label": "Good",
        "score": 1,
        "explanation": "Clear answer — you covered the main idea.",
        "ideal": "Normalization reduces duplication.",
    }
    report = client.post(
        "/api/viva/report",
        json={
            "subject": "DBMS",
            "topic": "Normalization",
            "level": "grad",
            "answers": [{**answer, "feedback": evaluated.json()}],
        },
    )
    assert report.status_code == 200
    assert set(report.json()) == {
        "id", "subject", "topic", "level", "score", "total",
        "strengths", "improvements", "questions", "createdAt",
    }
    assert report.json()["score"] == 1


def test_code_actions_and_basic_python_sandbox(client, monkeypatch):
    from app.routers import code

    async def fake_chat(messages, task="teacher", temperature=0.4):
        return "The loop stops before its exclusive boundary."

    monkeypatch.setattr(code, "chat", fake_chat)
    response = client.post(
        "/api/code",
        json={"action": "Explain", "code": "print(1)", "language": "Python", "level": "c9-10"},
    )
    assert response.status_code == 200
    assert response.json() == {"markdown": "The loop stops before its exclusive boundary."}
    for action in ("Debug", "Predict Output", "Give Hint", "Interview"):
        action_response = client.post(
            "/api/code",
            json={
                "action": action,
                "code": "print(1)",
                "language": "Python",
                "level": "c9-10",
                "hintIndex": 1,
            },
        )
        assert action_response.status_code == 200
        assert "markdown" in action_response.json()

    async def fake_json_call(messages, response_model, task="json"):
        if response_model is code.TestCaseSet:
            return response_model(
                markdown="Boundary tests included.",
                tests=[{"input": "n=1", "expected": "1", "actual": "0", "passed": False}],
            )
        return response_model(
            markdown="Generated exercise.",
            exercise={
                "id": "exercise-1",
                "title": "Loop practice",
                "topic": "loops",
                "difficulty": "Easy",
                "statement": "Sum values from 1 through n.",
                "examples": [{"input": "n=3", "output": "6"}],
                "starter": {"Python": "n=3"},
                "solution": "print(6)",
                "expectedOutput": "6",
                "hints": ["Hint 1", "Hint 2", "Hint 3"],
            },
        )

    monkeypatch.setattr(code, "json_call", fake_json_call)
    test_cases = client.post(
        "/api/code",
        json={
            "action": "Generate Test Cases",
            "code": "print(1)",
            "language": "Python",
            "level": "c9-10",
        },
    )
    assert test_cases.status_code == 200
    assert set(test_cases.json()) == {"markdown", "tests"}
    exercise = client.post(
        "/api/code",
        json={
            "action": "Generate Exercise",
            "code": "",
            "language": "Python",
            "level": "c9-10",
            "topic": "loops",
            "difficulty": "Easy",
        },
    )
    assert exercise.status_code == 200
    assert set(exercise.json()) == {"markdown", "exercise"}
    unsupported = client.post(
        "/api/code",
        json={"action": "Run", "code": "print(1)", "language": "Java", "level": "c9-10"},
    )
    assert "not executed" in unsupported.json()["markdown"]
    output, return_code = code.run_python("print(input().upper())", "hello")
    assert (output.strip(), return_code) == ("HELLO", 0)
    denied_output, denied_code = code.run_python("import socket\nsocket.socket()")
    assert denied_code != 0
    assert "unavailable" in denied_output


def test_progress_and_dashboard_empty_shapes_and_shared_mastery(client):
    progress = client.get("/api/progress")
    assert progress.status_code == 200
    snapshot = progress.json()
    assert set(snapshot) == {
        "mastery", "streak", "longestStreak", "questionsAttempted", "averageScore",
        "weeklyMinutes", "activity", "weakTopics", "studyPlan",
    }
    assert snapshot["mastery"] == {}
    assert snapshot["streak"] == 0
    assert len(snapshot["activity"]) == 84
    assert len(snapshot["studyPlan"]) == 7
    dashboard = client.get("/api/dashboard").json()
    assert dashboard["mastery"] == snapshot["mastery"]
    assert set(dashboard) == {"stats", "mastery", "suggestion", "activity", "plan"}


def test_practice_diagram_and_ppt_contracts(client, monkeypatch, caplog):
    from app.routers import diagram, ppt, practice

    section_calls: dict[str, int] = {}
    section_tasks: set[str] = set()
    section_token_budgets: list[int] = []

    async def fake_json_call(messages, response_model, task="json", **kwargs):
        if getattr(response_model, "__name__", "").startswith("GeneratedPracticeQuestionSet"):
            section_tasks.add(task)
            section_token_budgets.append(kwargs["max_tokens"])
            prompt = messages[-1]["content"]
            match = re.search(
                r"Create (\d+) (MCQ|short|long|numerical) questions, "
                r"each exactly (\d+) marks",
                prompt,
            )
            assert match is not None
            count, question_type, marks = match.groups()
            question_type = question_type.lower()
            section_calls[question_type] = section_calls.get(question_type, 0) + 1
            if question_type == "mcq" and section_calls[question_type] == 1:
                count = str(max(0, int(count) - 1))
            return response_model(
                questions=[
                    {
                        "type": question_type,
                        "text": f"{question_type} question {index}",
                        "marks": int(marks),
                        "options": ["A", "B", "C", "D"] if question_type == "mcq" else None,
                        "correct": 0 if question_type == "mcq" else None,
                        "model": "Expected answer",
                        "topic": "Electricity",
                        "keywords": ["current"],
                    }
                    for index in range(int(count))
                ]
            )
        if response_model is diagram.DiagramOutput:
            return response_model(
                title="Even or odd", code="flowchart TD\nA([Start]) --> B[End]", explanation="Follow the steps."
            )
        if response_model is ppt.DeckOutput:
            return response_model(
                slides=[
                    {"title": "Topic", "bullets": ["Intro"], "notes": "Introduce.", "kind": "title"},
                    {"title": "Idea", "bullets": ["One"], "notes": "Explain.", "kind": "content"},
                    {"title": "Practice", "bullets": ["Try it"], "notes": "Ask.", "kind": "content"},
                    {"title": "Review", "bullets": ["Recall"], "notes": "Review.", "kind": "content"},
                    {"title": "Conclusion", "bullets": ["Thanks"], "notes": "Close.", "kind": "conclusion"},
                ]
            )
        if response_model is practice.EvaluationSet:
            return response_model(
                evaluations=[
                    {
                        "id": "eval-1", "question": "What is Ohm's law?", "marks": 5,
                        "answer": "V = IR",
                        "rubric": [
                            {"label": "Concept", "got": 2, "max": 2},
                            {"label": "Explanation", "got": 1.25, "max": 1.25},
                            {"label": "Example", "got": 1, "max": 1},
                            {"label": "Presentation", "got": 0.75, "max": 0.75},
                        ],
                        "good": "Correct formula.", "improve": "Add an example.",
                        "model": "V = IR at constant temperature.", "topic": "Ohm's law",
                    }
                ]
            )
        if response_model is practice.Evaluation:
            return response_model(
                id="eval-1", question="What is Ohm's law?", marks=5, answer="V=IR with example",
                rubric=[
                    {"label": "Concept", "got": 2, "max": 2},
                    {"label": "Explanation", "got": 1.25, "max": 1.25},
                    {"label": "Example", "got": 1, "max": 1},
                    {"label": "Presentation", "got": 0.75, "max": 0.75},
                ],
                good="Good.", improve="Better.", model="V=IR.", topic="Ohm's law",
            )
        raise AssertionError(f"Unexpected response model {response_model}")

    monkeypatch.setattr(practice, "json_call", fake_json_call)
    from app.services import grading

    async def fake_rubric_call(messages, response_model, task="json", **kwargs):
        return response_model(
            rubric=[
                {"label": "Concept", "got": 0.8, "max": 0.8},
                {"label": "Explanation", "got": 0.5, "max": 0.5},
                {"label": "Example", "got": 0.4, "max": 0.4},
                {"label": "Presentation", "got": 0.3, "max": 0.3},
            ],
            feedback="Good start.",
        )

    monkeypatch.setattr(grading, "json_call", fake_rubric_call)
    async def fake_extract(data, filename, content_type):
        return [DocumentPage(1, "Question 1. State Ohm's law.\nAnswer: V = IR.")]

    monkeypatch.setattr(practice, "extract_document", fake_extract)
    monkeypatch.setattr(diagram, "json_call", fake_json_call)
    monkeypatch.setattr(ppt, "json_call", fake_json_call)
    generated = client.post(
        "/api/practice/generate",
        json={
            "level": "c9-10", "subject": "Science", "chapter": "Electricity",
            "difficulty": "Medium", "totalMarks": 20, "timeMin": 40, "weakFocus": [],
        },
    )
    assert generated.status_code == 200, generated.text
    paper = generated.json()
    assert set(paper) == {
        "id", "title", "subject", "level", "chapter", "difficulty", "totalMarks",
        "timeMin", "createdAt", "weakFocus", "sections", "questions",
    }
    assert sum(question["marks"] for question in paper["questions"]) == 20
    assert sum(section["marks"] for section in paper["sections"]) == 20
    assert {section["name"] for section in paper["sections"]} == {
        "MCQ", "short", "long", "numerical"
    }
    assert section_calls == {"mcq": 3, "short": 2, "long": 1, "numerical": 1}
    assert section_tasks == {"practice"}
    assert set(section_token_budgets) == {1920, 480, 680, 340, 500, 440}
    for section_type in ("MCQ", "short", "long", "numerical"):
        assert f"practice_generation section={section_type}" in caplog.text
        assert "duration_ms=" in caplog.text
    submitted = client.post(
        f"/api/practice/{paper['id']}/submit",
        json={"answers": {"q-A-1": "0", "q-B-1": "It carries current."}, "timeUsedSec": 240},
    )
    assert submitted.status_code == 200
    assert set(submitted.json()) == {
        "score", "total", "timeUsedSec", "perQ", "sections", "weakTopics"
    }
    assert submitted.json()["perQ"][0]["status"] == "correct"
    short_result = next(
        item for item in submitted.json()["perQ"] if item["id"] == "q-B-1"
    )
    assert short_result["awarded"] == 2
    checked = client.post(
        "/api/practice/check",
        files={"answers": ("answers.pdf", b"%PDF-1.7 mock", "application/pdf")},
    )
    assert checked.status_code == 200
    rechecked = client.post("/api/practice/eval-1/re-evaluate", json={"answer": "V=IR with example"})
    assert rechecked.status_code == 200
    generated_diagram = client.post(
        "/api/diagram",
        json={"prompt": "Even or odd", "type": "Flowchart", "level": "c9-10"},
    )
    assert generated_diagram.status_code == 200
    assert set(generated_diagram.json()) == {
        "id", "prompt", "title", "type", "code", "explanation", "level", "createdAt"
    }
    generated_deck = client.post(
        "/api/ppt",
        json={
            "topic": "Artificial Intelligence", "level": "c9-10", "slides": 5,
            "theme": "Indigo Modern", "speakerNotes": True, "includeDiagrams": False,
            "includeQuiz": False, "extra": "",
        },
    )
    assert generated_deck.status_code == 200
    assert set(generated_deck.json()) == {
        "id", "topic", "level", "theme", "slides", "createdAt", "speakerNotes"
    }
    assert len(generated_deck.json()["slides"]) == 5


def test_truncated_practice_batch_splits_without_repeating_full_request(client, monkeypatch):
    from app.routers import practice

    batch_sizes: list[int] = []

    async def fake_json_call(messages, response_model, task="json", **kwargs):
        prompt = messages[-1]["content"]
        match = re.search(
            r"Create (\d+) (MCQ|short|long|numerical) questions, "
            r"each exactly (\d+) marks",
            prompt,
        )
        assert match is not None
        count, question_type, marks = match.groups()
        count = int(count)
        question_type = question_type.lower()
        if question_type == "mcq":
            batch_sizes.append(count)
            if count == 4:
                raise practice.CompletionTruncated('{"questions":[', 100)
        return response_model(
            questions=[
                {
                    "type": question_type,
                    "text": f"{question_type} question {index}",
                    "marks": int(marks),
                    "options": ["A", "B", "C", "D"] if question_type == "mcq" else None,
                    "correct": 0 if question_type == "mcq" else None,
                    "model": "Expected answer",
                    "topic": "Electricity",
                    "keywords": ["current"],
                }
                for index in range(count)
            ]
        )

    monkeypatch.setattr(practice, "json_call", fake_json_call)
    response = client.post(
        "/api/practice/generate",
        json={
            "level": "c9-10",
            "subject": "Science",
            "chapter": "Electricity",
            "difficulty": "Medium",
            "totalMarks": 20,
            "timeMin": 40,
            "weakFocus": [],
        },
    )

    assert response.status_code == 200, response.text
    assert batch_sizes == [4, 1, 2, 2]
    assert len(response.json()["questions"]) == 10


def test_practice_parser_normalizes_model_question_type_variants():
    from app.routers.practice import GeneratedMcqQuestion, _valid_partial_questions

    raw = (
        '{"questions":['
        '{"type":"MCQ","text":"Question?","marks":1,'
        '"options":["A","B"],"correct":0,"model":"Answer",'
        '"topic":"Topic","keywords":[]}'
        "]"
    )

    questions = _valid_partial_questions(raw, GeneratedMcqQuestion, "MCQ", 1)

    assert len(questions) == 1
    assert questions[0].type == "mcq"


def test_practice_generation_returns_reason_for_mocked_upstream_failure(client, monkeypatch):
    from app.routers import practice

    calls = 0

    async def failed_json_call(*args, **kwargs):
        nonlocal calls
        calls += 1
        raise OmniRouteError(
            "OmniRoute returned HTTP 429: rate limited",
            status_code=429,
            error_type="RateLimitError",
            response_excerpt="rate limited",
        )

    monkeypatch.setattr(practice, "json_call", failed_json_call)
    response = client.post(
        "/api/practice/generate",
        json={
            "level": "c9-10",
            "subject": "Science",
            "chapter": "Electricity",
            "difficulty": "Medium",
            "totalMarks": 20,
            "timeMin": 40,
            "weakFocus": [],
        },
    )

    assert response.status_code == 502
    assert "section A (MCQ) failed" in response.json()["detail"]
    assert "HTTP 429" in response.json()["detail"]
    assert calls == 6


def test_pptx_export_route_returns_openable_powerpoint_with_notes(client):
    from io import BytesIO

    from pptx import Presentation

    response = client.post(
        "/api/pptx",
        json={
            "topic": "Fractions",
            "theme": "Playful",
            "speakerNotes": True,
            "slides": [
                {
                    "id": "s1",
                    "title": "Fractions",
                    "bullets": ["Parts of a whole"],
                    "notes": "Introduce the idea.",
                    "kind": "title",
                },
                {
                    "id": "s2",
                    "title": "Equivalent fractions",
                    "bullets": ["Multiply the top and bottom by the same number."],
                    "notes": "Show one half equals two fourths.",
                    "kind": "conclusion",
                },
            ],
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    deck = Presentation(BytesIO(response.content))
    assert len(deck.slides) == 2
    assert "Introduce the idea." in deck.slides[0].notes_slide.notes_text_frame.text


def test_upload_validation_and_error_contracts(client):
    bad_type = client.post(
        "/api/docs/upload",
        files={"file": ("document.exe", b"not a document", "application/octet-stream")},
    )
    assert bad_type.status_code == 415
    assert set(bad_type.json()) == {"error", "detail"}
    oversized = client.post(
        "/api/docs/upload",
        files={"file": ("large.pdf", b"%PDF-" + b"x" * (20 * 1024 * 1024 + 1), "application/pdf")},
    )
    assert oversized.status_code == 413


def test_reset_progress_clears_learning_activity_and_xp(client):
    db_generator = app.dependency_overrides[get_db]()
    db = next(db_generator)
    try:
        db.add(ChatSession(id="reset-chat"))
        db.add(
            ChatMessage(
                session_id="reset-chat",
                role="user",
                content="Previously asked question",
            )
        )
        db.add(
            PracticePaper(
                id="reset-paper",
                paper_json=json.dumps({"questions": []}),
            )
        )
        db.add(
            PracticeAttempt(
                paper_id="reset-paper",
                subject="Science",
                chapter="Electricity",
                total=1,
                score=1,
                paper_json=json.dumps(
                    {"questions": [{"id": "q1", "topic": "Current", "marks": 1}]}
                ),
                result_json=json.dumps(
                    {"perQ": [{"id": "q1", "awarded": 1}], "timeUsedSec": 60}
                ),
            )
        )
        db.add(
            PracticeEvaluation(
                id="reset-evaluation",
                evaluation_json=json.dumps({"question": "Previously checked answer"}),
            )
        )
        db.add(
            VivaSession(
                id="reset-viva",
                subject="Science",
                topic="Gravity",
                level="c9-10",
                questions_json="[]",
                answers_json="[]",
                report_json=json.dumps(
                    {
                        "total": 1,
                        "score": 1,
                        "questions": [
                            {
                                "topic": "Gravity",
                                "feedback": {"score": 1},
                            }
                        ],
                    }
                ),
            )
        )
        db.commit()
    finally:
        db_generator.close()

    response = client.delete("/api/progress")
    assert response.status_code == 200
    assert response.json() == {"status": "cleared"}

    dashboard = client.get("/api/dashboard")
    assert dashboard.status_code == 200
    assert dashboard.json()["stats"] == {
        "dailyProgress": 0,
        "streak": 0,
        "xpWeek": 0,
        "solved": 0,
    }
