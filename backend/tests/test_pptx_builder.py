from io import BytesIO

import pytest
from pptx import Presentation

from app.services.pptx_builder import build_pptx


@pytest.mark.parametrize(
    "theme",
    ["Indigo Modern", "Clean White", "Dark Elegant", "Playful"],
)
def test_build_pptx_reopens_with_slide_content_and_speaker_notes(theme):
    data = build_pptx(
        topic="Fractions",
        theme=theme,
        slides=[
            {
                "title": "Fractions",
                "bullets": ["Parts of a whole"],
                "notes": "Introduce the lesson.",
                "kind": "title",
            },
            {
                "title": "Equivalent fractions",
                "bullets": ["Multiply numerator and denominator by the same number."],
                "notes": "Use one half and two fourths as an example.",
                "kind": "content",
            },
        ],
        speaker_notes=True,
    )

    deck = Presentation(BytesIO(data))
    assert len(deck.slides) == 2
    assert "Fractions" in deck.slides[0].shapes.title.text if deck.slides[0].shapes.title else any(
        "Fractions" in shape.text for shape in deck.slides[0].shapes if shape.has_text_frame
    )
    assert "Introduce the lesson." in deck.slides[0].notes_slide.notes_text_frame.text
    assert "Use one half" in deck.slides[1].notes_slide.notes_text_frame.text


def test_build_pptx_can_omit_speaker_notes():
    data = build_pptx(
        topic="Cells",
        theme="Clean White",
        slides=[{"title": "Cells", "bullets": [], "notes": "Do not include.", "kind": "title"}],
        speaker_notes=False,
    )

    deck = Presentation(BytesIO(data))
    assert len(deck.slides) == 1
    assert "Do not include." not in deck.slides[0].notes_slide.notes_text_frame.text
